"""Synthetic-only tests; no coursework input, network, models, or persistent outputs.

Run with python3.13 -I -B build_berkay_dataset.py --self-test, or run this file
with python3.13 -I -B. Every input/output fixture lives in TemporaryDirectory.
"""

from __future__ import annotations

from collections import Counter
from copy import deepcopy
import importlib.util
import io
from pathlib import Path
import struct
import sys
import tempfile
import unittest
from unittest import mock
import warnings
import zipfile

from openpyxl import load_workbook


# Explicit sibling-file loading works with -I without adding a directory to
# sys.path. Disable bytecode writes even when a caller forgets -B.
sys.dont_write_bytecode = True
if "BUILDER" not in globals():
    spec = importlib.util.spec_from_file_location(
        "berkay_builder_under_test", Path(__file__).with_name("build_berkay_dataset.py")
    )
    BUILDER = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = BUILDER
    spec.loader.exec_module(BUILDER)
B = BUILDER


def synthetic_inputs(count: int = 1) -> tuple[dict, dict]:
    """Produce distinct fake full contexts and exact native multi-span annotations."""
    source = {"version": "v1.0", "data": []}
    curation = {
        "schema_version": "1.0", "created_date": "2026-10-13",
        "source": {"filename": "CUADv1.json", "sha256": "", "version": "v1.0"},
        "owner": "Berkay", "reviewer": "Engin", "review_status": "pending",
        "expected_contract_count": count, "contracts": [],
    }
    for i in range(count):
        title = f"Synthetic Agreement {i}"
        renewal = "Renewal: ten (10) years, with no extension."
        law = "Law: the State of Illinois governs this agreement."
        footer = "Footer unchanged."
        context = f" \r\nSynthetic original {i} — café.\r\nHeader.\n{renewal}\r\nSpacer clause stays.\n{law}\n{footer}"

        def span(quote: str) -> dict:
            start = context.index(quote)
            return {"start": start, "end": start + len(quote), "text": quote}

        def patch(before: str, after: str) -> dict:
            start = context.index(before)
            return {"start": start, "end": start + len(before), "before": before, "after": after}

        source["data"].append({"title": title, "paragraphs": [{"context": context, "qas": [
            {"id": f"{title}__Renewal Term", "question": "Extract renewal text.", "is_impossible": False,
             "answers": [{"answer_start": context.index(quote), "text": quote} for quote in (renewal, footer)]},
            {"id": f"{title}__Governing Law", "question": "Extract law text.", "is_impossible": False,
             "answers": [{"answer_start": context.index(law), "text": law}]},
        ]}]})
        curation["contracts"].append({
            "id": f"BER-{i + 1:02d}-Synthetic", "original_contract_id": title,
            "contract_type": "Synthetic Research Agreement", "document_index": i, "paragraph_index": 0,
            "context_sha256": B.digest(context.encode("utf-8")),
            "questions": [
                {"id": "Q1", "category": "Renewal Term", "text": "How long is the renewal?",
                 "answer_original": "Ten years.", "answer_modified": "Fifteen years.",
                 "annotation_id": f"{title}__Renewal Term",
                 "evidence": [span("ten (10) years, with no extension."), span(footer)],
                 "scope_notes": "Only the duration of this synthetic renewal."},
                {"id": "Q2", "category": "Governing Law", "text": "Which state governs?",
                 "answer_original": "Illinois.", "answer_modified": "Massachusetts.",
                 "annotation_id": f"{title}__Governing Law", "evidence": [span("State of Illinois")]},
            ],
            "variants": {
                "B1": {"affected_question": "Q1", "semantic_change": "Synthetic renewal duration 10 to 15 years.",
                       "patches": [patch("10", "15"), patch("ten", "fifteen")]},
                "B2": {"affected_question": "Q2", "semantic_change": "Synthetic governing state only.",
                       "patches": [patch("Illinois", "Massachusetts")]},
            },
            "independence_note": "Separated toy clauses; no claim about real legal independence.",
            "review_notes": [],
        })
    curation["source"]["sha256"] = B.digest(B.json_bytes(source))
    return source, curation


class BuilderTests(unittest.TestCase):
    """Positive and fail-closed negative checks on synthetic data only."""

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="berkay-synthetic-test-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.input_dir = self.root / "input"
        self.input_dir.mkdir()
        self.source_path = self.input_dir / "CUADv1.json"
        self.curation_path = self.input_dir / "curation.json"
        self.output = self.root / "synthetic-dataset"
        self.archive = self.root / "synthetic-dataset.zip"
        self.source, self.curation = synthetic_inputs()

    @property
    def contract(self) -> dict:
        return self.curation["contracts"][0]

    def write_inputs(self, *, refresh_hash: bool = True) -> None:
        source_raw = B.json_bytes(self.source)
        if refresh_hash:
            self.curation["source"]["sha256"] = B.digest(source_raw)
        self.source_path.write_bytes(source_raw)
        self.curation_path.write_bytes(B.json_bytes(self.curation))

    def prepare(self, *, refresh_hash: bool = True):
        self.write_inputs(refresh_hash=refresh_hash)
        return B.prepare(self.source_path.read_bytes(), self.curation_path.read_bytes(), self.source_path.name)

    def build(self):
        self.write_inputs()
        return B.build(self.source_path, self.curation_path, self.output, self.archive)

    def assert_prepare_fails(self, message: str, *, refresh_hash: bool = True) -> None:
        with self.assertRaisesRegex(B.ValidationError, message):
            self.prepare(refresh_hash=refresh_hash)

    def test_ten_contract_package_has_exact_counts_bytes_rows_and_archive(self) -> None:
        self.source, self.curation = synthetic_inputs(10)
        prepared = self.build()
        self.assertEqual(len(prepared.documents), 40)
        self.assertEqual(len(prepared.rows), 60)
        self.assertEqual(Counter(r["case_label"] for r in prepared.rows), Counter({label: 20 for label in B.LABELS}))
        self.assertEqual(len(prepared.expected_files), 45)
        self.assertEqual((self.output / "curation.json").read_bytes(), self.curation_path.read_bytes())
        for i, contract in enumerate(self.curation["contracts"]):
            original = self.source["data"][i]["paragraphs"][0]["context"].encode("utf-8")
            prefix = self.output / "documents" / contract["id"]
            self.assertEqual((prefix / "A.txt").read_bytes(), original)
            self.assertEqual((prefix / "B0.txt").read_bytes(), original)
            self.assertFalse(original.endswith(b"\n"))
            self.assertIn(b"\r\n", original)
        B.validate_package(self.output, prepared)
        B.validate_zip(self.archive, self.output, prepared)
        B.check_unchanged(self.source_path, self.curation_path, prepared)
        with zipfile.ZipFile(self.archive) as archive:
            self.assertEqual(len(archive.namelist()), 45)
            self.assertTrue(all(name.startswith(self.output.name + "/") for name in archive.namelist()))
            # Extraction into a fresh temp folder also proves relative links resolve.
            extracted = self.root / "extracted"
            extracted.mkdir()
            archive.extractall(extracted)
            B.validate_package(extracted / self.output.name, prepared)
        report = B.parse_json((self.output / "validation.json").read_bytes(), "test report")
        self.assertFalse(report["semantic_or_legal_approval"])
        self.assertFalse(report["inference_performed"])
        self.assertEqual(report["human_review_status"], "pending")
        self.assertNotIn("validation.json", report["sha256"])
        self.assertNotIn(self.archive.name, report["sha256"])
        self.assertNotIn(str(self.root), (self.output / "validation.json").read_text())

    def test_patches_sorted_and_variants_independent_with_offset_remapping(self) -> None:
        prepared = self.prepare()
        prefix = f"documents/{self.contract['id']}"
        original = prepared.documents[f"{prefix}/A.txt"].decode("utf-8")
        b1 = prepared.documents[f"{prefix}/B1.txt"].decode("utf-8")
        b2 = prepared.documents[f"{prefix}/B2.txt"].decode("utf-8")
        self.assertEqual(b1, original.replace("ten (10)", "fifteen (15)"))
        self.assertEqual(b2, original.replace("Illinois", "Massachusetts"))
        self.assertNotIn("fifteen", b2)
        for row in prepared.rows:
            a_spans, b_spans = row["evidence"]["A"], row["evidence"]["B"]
            if row["variant"] == "B1" and row["question_id"] == "Q2":
                self.assertEqual(b_spans[0]["start"], a_spans[0]["start"] + 4)
                self.assertEqual(b_spans[0]["text"], a_spans[0]["text"])
            if row["variant"] == "B2" and row["question_id"] == "Q1":
                self.assertEqual(b_spans[0], a_spans[0])
                self.assertEqual(b_spans[1]["start"], a_spans[1]["start"] + 5)
                self.assertEqual(b_spans[1]["text"], a_spans[1]["text"])
            self.assertEqual(row["answer_supported_by_B"],
                             self.contract["questions"][int(row["question_id"][1]) - 1]
                             ["answer_modified" if row["case_label"] == "relevant_conflict" else "answer_original"])
            if row["variant"] == "B0":
                self.assertIn("No edits", row["exactly_what_changed"])
            elif row["variant"] == "B1":
                self.assertIn('"ten" -> "fifteen"', row["exactly_what_changed"])
                self.assertNotIn("+30", row["exactly_what_changed"])

    def test_optional_scope_and_review_notes_can_be_absent_or_empty(self) -> None:
        self.contract.pop("review_notes")
        for question in self.contract["questions"]:
            question.pop("scope_notes", None)
        self.prepare()
        self.contract["review_notes"] = None
        self.contract["questions"][0]["scope_notes"] = ""
        self.prepare()

    def test_stale_source_hash(self) -> None:
        self.curation["source"]["sha256"] = "0" * 64
        self.assert_prepare_fails("stale SHA-256", refresh_hash=False)

    def test_stale_context_hash(self) -> None:
        self.contract["context_sha256"] = "1" * 64
        self.assert_prepare_fails("stale SHA-256")

    def test_wrong_version_and_filename(self) -> None:
        self.curation["source"]["version"] = "1.0"
        self.assert_prepare_fails("version mismatch")
        self.curation["source"]["version"] = "v1.0"
        self.curation["source"]["filename"] = "different.json"
        self.assert_prepare_fails("filename does not match")
        self.curation["source"]["filename"] = "/personal/CUADv1.json"
        self.assert_prepare_fails("portable ASCII basename")

    def test_title_document_index_and_single_full_context_enforced(self) -> None:
        original = deepcopy(self.contract)
        for name, value, error in (
            ("original_contract_id", "Wrong title", "title/index"),
            ("document_index", 10, "out of range"),
            ("document_index", True, "nonnegative integer"),
            ("paragraph_index", 1, "exactly one full context"),
        ):
            with self.subTest(name=name, value=value):
                self.contract[name] = value
                self.assert_prepare_fails(error)
                self.curation["contracts"][0] = deepcopy(original)
        self.source["data"][0]["paragraphs"].append(deepcopy(self.source["data"][0]["paragraphs"][0]))
        self.assert_prepare_fails("exactly one full context")

    def test_stale_patch_offsets_before_and_noop_rejected(self) -> None:
        original = deepcopy(self.contract["variants"]["B1"]["patches"])
        for name, value, error in (
            ("start", 0, "Stale patch"), ("end", 999999, "out of range"),
            ("before", "99", "Stale patch"), ("before", "", "empty text"),
            ("after", "10", "No-op patch"),
        ):
            with self.subTest(name=name, value=value):
                self.contract["variants"]["B1"]["patches"] = deepcopy(original)
                self.contract["variants"]["B1"]["patches"][0][name] = value
                self.assert_prepare_fails(error)

    def test_empty_null_and_wrong_target_variants_rejected(self) -> None:
        original = deepcopy(self.contract["variants"]["B1"])
        for value, error in (({}, "missing fields"), (None, "expected an object"),
                             ({**original, "patches": []}, "nonempty patches"),
                             ({**original, "affected_question": "Q2"}, "must affect Q1")):
            with self.subTest(value=value):
                self.contract["variants"]["B1"] = value
                self.assert_prepare_fails(error)

    def test_overlapping_patches_rejected(self) -> None:
        self.contract["variants"]["B1"]["patches"].append(deepcopy(self.contract["variants"]["B1"]["patches"][0]))
        self.assert_prepare_fails("Overlapping patches")

    def test_patch_boundary_crossing_and_outside_target_rejected(self) -> None:
        context = self.source["data"][0]["paragraphs"][0]["context"]
        start = self.contract["questions"][0]["evidence"][0]["start"]
        for a, b in ((start - 1, start + 2), (start - 2, start - 1)):
            with self.subTest(start=a):
                self.contract["variants"]["B1"]["patches"] = [
                    {"start": a, "end": b, "before": context[a:b], "after": "changed"}
                ]
                self.assert_prepare_fails("wholly inside target evidence")

    def test_b1_cannot_edit_q2_evidence(self) -> None:
        self.contract["variants"]["B1"]["patches"] = deepcopy(self.contract["variants"]["B2"]["patches"])
        self.assert_prepare_fails("wholly inside target evidence")

    def test_missing_wrong_and_stale_evidence_rejected(self) -> None:
        original = deepcopy(self.contract["questions"][0]["evidence"])
        for value, error in (([], "Missing evidence"),
                             ([{**original[0], "text": "wrong"}], "wrong evidence text"),
                             ([{**original[0], "start": original[0]["start"] + 1}], "stale offset")):
            with self.subTest(value=value):
                self.contract["questions"][0]["evidence"] = value
                self.assert_prepare_fails(error)

    def test_evidence_outside_selected_native_answer_rejected(self) -> None:
        context = self.source["data"][0]["paragraphs"][0]["context"]
        quote = "Spacer clause stays."
        start = context.index(quote)
        self.contract["questions"][0]["evidence"] = [{"start": start, "end": start + len(quote), "text": quote}]
        self.assert_prepare_fails("not a subspan")

    def test_overlap_between_questions_fails_closed(self) -> None:
        self.contract["questions"][1]["category"] = "Renewal Term"
        self.contract["questions"][1]["annotation_id"] = self.contract["questions"][0]["annotation_id"]
        self.contract["questions"][1]["evidence"] = deepcopy(self.contract["questions"][0]["evidence"])
        self.assert_prepare_fails("Overlapping evidence")

    def test_overlap_within_one_question_fails_closed(self) -> None:
        self.contract["questions"][0]["evidence"].append(deepcopy(self.contract["questions"][0]["evidence"][0]))
        self.assert_prepare_fails("Overlapping evidence")

    def test_annotation_id_category_and_exact_native_text_checked(self) -> None:
        question = self.contract["questions"][0]
        original_id = question["annotation_id"]
        question["annotation_id"] = "missing__Renewal Term"
        self.assert_prepare_fails("Missing or ambiguous annotation_id")
        question["annotation_id"] = original_id
        question["category"] = "Wrong Category"
        self.assert_prepare_fails("Unsupported clause category")
        question["category"] = "Renewal Term"
        native = self.source["data"][0]["paragraphs"][0]["qas"][0]
        native["answers"][0]["text"] = native["answers"][0]["text"].replace("ten", "TEN")
        self.assert_prepare_fails("Native answer text/offset")

    def test_disallowed_real_cuad_category_is_rejected(self) -> None:
        question = self.contract["questions"][0]
        question["category"] = "Parties"
        question["annotation_id"] = self.contract["original_contract_id"] + "__Parties"
        self.source["data"][0]["paragraphs"][0]["qas"][0]["id"] = question["annotation_id"]
        self.assert_prepare_fails("Unsupported clause category")

    def test_duplicate_annotation_ids_rejected(self) -> None:
        qas = self.source["data"][0]["paragraphs"][0]["qas"]
        qas.append(deepcopy(qas[0]))
        self.assert_prepare_fails("Missing or ambiguous annotation_id")

    def test_duplicate_contracts_and_contexts_rejected(self) -> None:
        self.curation["expected_contract_count"] = 2
        self.curation["contracts"].append(deepcopy(self.contract))
        self.assert_prepare_fails("Duplicate contract")
        self.source, self.curation = synthetic_inputs(2)
        context = self.source["data"][0]["paragraphs"][0]["context"]
        self.source["data"][1]["paragraphs"][0]["context"] = context
        self.curation["contracts"][1]["context_sha256"] = B.digest(context.encode())
        self.assert_prepare_fails("Duplicate full contract context")

    def test_casefold_contract_ids_rejected(self) -> None:
        self.source, self.curation = synthetic_inputs(2)
        self.curation["contracts"][1]["id"] = self.contract["id"].lower()
        self.assert_prepare_fails("Duplicate contract")

    def test_wrong_expected_count_and_question_count_rejected(self) -> None:
        self.curation["expected_contract_count"] = 10
        self.assert_prepare_fails("Contract count")
        self.curation["expected_contract_count"] = 1
        self.contract["questions"].pop()
        self.assert_prepare_fails("exactly Q1 and Q2")

    def test_inconsistent_answers_and_fake_approval_rejected(self) -> None:
        question = self.contract["questions"][0]
        original = question["answer_modified"]
        for answer in ("", question["answer_original"], "  " + question["answer_original"] + " "):
            with self.subTest(answer=answer):
                question["answer_modified"] = answer
                self.assert_prepare_fails("empty text|Inconsistent answers")
        question["answer_modified"] = original
        self.curation["review_status"] = "approved"
        self.assert_prepare_fails("must remain pending")

    def test_answer_paraphrases_are_not_claimed_as_semantically_validated(self) -> None:
        # Intentionally unverifiable paraphrase: mechanical checker must not pretend
        # it can judge its meaning. Pending review is the explicit trust boundary.
        self.contract["questions"][0]["answer_modified"] = "A curator-provided paraphrase requiring review."
        prepared = self.prepare()
        self.assertTrue(all(row["review_status"] == "pending" for row in prepared.rows))
        self.assertIn("semantic review pending", prepared.rows[0]["provenance"]["answer_origin"])

    def test_wrong_row_counts_labels_answers_and_matched_pairs_rejected(self) -> None:
        prepared = self.prepare()
        with self.assertRaisesRegex(B.ValidationError, "six rows"):
            B.validate_rows(prepared.rows[:-1], 1)
        for column, value, index, error in (
            ("case_label", "agreement", 2, "Incorrect case label"),
            ("answer_supported_by_B", "Fifteen years.", 0, "Inconsistent source answers"),
            ("review_status", "approved", 0, "pending"),
            ("expected_behaviour", "ANSWER", 2, "expected behaviour"),
        ):
            with self.subTest(column=column):
                rows = deepcopy(prepared.rows)
                rows[index][column] = value
                with self.assertRaisesRegex(B.ValidationError, error):
                    B.validate_rows(rows, 1)
        rows = deepcopy(prepared.rows)
        rows[1] = deepcopy(rows[0])
        rows[1]["example_id"] += "-not-a-duplicate-id"
        with self.assertRaisesRegex(B.ValidationError, "matched question/variant"):
            B.validate_rows(rows, 1)

    def test_unsafe_paths_and_reserved_windows_names_rejected(self) -> None:
        for key in ("../escape", "a/b", "a\\b", "C:drive", "CON", "nul.txt", "LPT9", "trailing.",
                    "percent%20name", "with space", ".hidden", "\x00bad"):
            with self.subTest(key=key):
                self.contract["id"] = key
                self.assert_prepare_fails("portable ASCII basename|reserved Windows")

    def test_formula_and_error_strings_are_stored_as_literal_text(self) -> None:
        self.contract["contract_type"] = "#N/A"
        self.contract["questions"][0]["text"] = '=HYPERLINK("https://example.invalid", "not executed")'
        self.contract["questions"][0]["answer_original"] = "+2+2"
        self.contract["questions"][0]["answer_modified"] = "@SUM(A1)"
        self.contract["questions"][1]["text"] = "-2+2"
        prepared = self.build()
        workbook = load_workbook(self.output / B.WORKBOOK, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook[B.SHEET]
        self.assertEqual(sheet["E2"].value, self.contract["questions"][0]["text"])
        self.assertEqual(sheet["C2"].value, "#N/A")
        self.assertTrue(all(cell.data_type == "s" for row in sheet for cell in row))
        B.validate_package(self.output, prepared)

    def test_overlength_or_invalid_spreadsheet_text_rejected_before_truncation(self) -> None:
        for value in ("a" * 32768, "\U0001F680" * 16384, "bad\x00text", "bad￿text"):
            with self.subTest(length=len(value)):
                self.contract["questions"][0]["text"] = value
                self.assert_prepare_fails("32767|XML-invalid")
        self.contract["questions"][0]["text"] = "a" * 32767
        self.prepare()

    def test_missing_tampered_and_unexpected_document_files_rejected(self) -> None:
        prepared = self.build()
        name = next(iter(prepared.documents))
        document = self.output / name
        original = document.read_bytes()
        document.unlink()
        with self.assertRaisesRegex(B.ValidationError, "Package file mismatch"):
            B.validate_package(self.output, prepared)
        document.write_bytes(original + b"tampered")
        with self.assertRaisesRegex(B.ValidationError, "tampered"):
            B.validate_package(self.output, prepared)
        document.write_bytes(original)
        (self.output / "unrelated.txt").write_text("not allowed")
        with self.assertRaisesRegex(B.ValidationError, "Package file mismatch"):
            B.validate_package(self.output, prepared)

    def test_symlink_inside_package_rejected_without_following(self) -> None:
        prepared = self.build()
        name = next(iter(prepared.documents))
        destination = self.output / name
        external = self.root / "unrelated.txt"
        external.write_text("untouched")
        destination.unlink()
        destination.symlink_to(external)
        with self.assertRaisesRegex(B.ValidationError, "Symlink"):
            B.validate_package(self.output, prepared)
        self.assertEqual(external.read_text(), "untouched")

    def test_b2_accidentally_derived_from_b1_detected(self) -> None:
        prepared = self.build()
        prefix = self.output / "documents" / self.contract["id"]
        bad_b2 = (prefix / "B1.txt").read_bytes().replace(b"Illinois", b"Massachusetts")
        (prefix / "B2.txt").write_bytes(bad_b2)
        with self.assertRaisesRegex(B.ValidationError, "incorrectly derived"):
            B.validate_package(self.output, prepared)

    def test_jsonl_tampered_quote_and_answer_detected(self) -> None:
        prepared = self.build()
        path = self.output / "examples.jsonl"
        original = path.read_bytes()
        for field, value in (("answer_supported_by_A", "tampered"), ("question", "different")):
            with self.subTest(field=field):
                rows = deepcopy(prepared.rows)
                rows[0][field] = value
                path.write_bytes(b"".join(B.json_bytes(row).replace(b"\n", b" ") + b"\n" for row in rows))
                with self.assertRaises(B.ValidationError):
                    B.validate_package(self.output, prepared)
        path.write_bytes(original)
        rows = deepcopy(prepared.rows)
        rows[0]["evidence"]["A"][0]["text"] = "wrong quote"
        path.write_bytes(b"".join(B.json_bytes(row).replace(b"\n", b" ") + b"\n" for row in rows))
        with self.assertRaisesRegex(B.ValidationError, "canonical rows"):
            B.validate_package(self.output, prepared)

    def test_workbook_mismatch_formula_and_broken_link_detected(self) -> None:
        prepared = self.build()
        path = self.output / B.WORKBOOK
        original = path.read_bytes()
        for kind in ("value", "formula", "link"):
            with self.subTest(kind=kind):
                path.write_bytes(original)
                workbook = load_workbook(path)
                sheet = workbook[B.SHEET]
                if kind == "value":
                    sheet["E2"] = "different question"
                elif kind == "formula":
                    sheet["E2"] = "=1+1"
                else:
                    sheet["F2"].hyperlink = "../escape.txt"
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(B.ValidationError, "workbook mismatch|hyperlink"):
                    B.validate_package(self.output, prepared)

    def test_validation_report_cannot_be_changed_into_human_approval(self) -> None:
        prepared = self.build()
        path = self.output / "validation.json"
        report = B.parse_json(path.read_bytes(), "report")
        report["semantic_or_legal_approval"] = True
        path.write_bytes(B.json_bytes(report))
        with self.assertRaisesRegex(B.ValidationError, "Validation report"):
            B.validate_package(self.output, prepared)

    def test_existing_destinations_are_never_overwritten(self) -> None:
        self.write_inputs()
        self.output.mkdir()
        external = self.output / "unrelated.txt"
        external.write_text("preserve this")
        with self.assertRaisesRegex(B.ValidationError, "Destination already exists"):
            B.build(self.source_path, self.curation_path, self.output, self.archive)
        self.assertEqual(external.read_text(), "preserve this")
        external.unlink()
        with self.assertRaisesRegex(B.ValidationError, "Destination already exists"):
            B.build(self.source_path, self.curation_path, self.output, self.archive)
        self.output.rmdir()
        self.archive.write_bytes(b"preserve archive")
        with self.assertRaisesRegex(B.ValidationError, "Destination already exists"):
            B.build(self.source_path, self.curation_path, self.output, self.archive)
        self.assertEqual(self.archive.read_bytes(), b"preserve archive")
        self.assertFalse(self.output.exists())

    def test_output_zip_collision_and_dangling_symlink_rejected(self) -> None:
        self.write_inputs()
        for archive in (self.output, self.root / self.output.name.upper()):
            with self.subTest(archive=archive):
                with self.assertRaisesRegex(B.ValidationError, "case-insensitive collisions"):
                    B.build(self.source_path, self.curation_path, self.output, archive)
        self.output.symlink_to(self.root / "nonexistent")
        with self.assertRaisesRegex(B.ValidationError, "Destination already exists"):
            B.build(self.source_path, self.curation_path, self.output, self.archive)

    def test_input_change_during_staging_stops_publication(self) -> None:
        self.write_inputs()
        writer = B.write_workbook

        def change_input(path, rows):
            writer(path, rows)
            self.source_path.write_bytes(self.source_path.read_bytes() + b"\n")

        with mock.patch.object(B, "write_workbook", side_effect=change_input):
            with self.assertRaisesRegex(B.ValidationError, "Source changed"):
                B.build(self.source_path, self.curation_path, self.output, self.archive)
        self.assertFalse(self.output.exists())
        self.assertFalse(self.archive.exists())
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ["input"])

    def test_late_zip_failure_retains_only_valid_output_not_final_zip(self) -> None:
        self.write_inputs()
        with mock.patch.object(B, "validate_zip", side_effect=B.ValidationError("synthetic archive failure")):
            with self.assertRaisesRegex(B.ValidationError, "synthetic archive failure"):
                B.build(self.source_path, self.curation_path, self.output, self.archive)
        self.assertTrue(self.output.is_dir())
        self.assertFalse(self.archive.exists())
        prepared = B.prepare(self.source_path.read_bytes(), self.curation_path.read_bytes(), self.source_path.name)
        B.validate_package(self.output, prepared)
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ["input", self.output.name])

    def test_zip_missing_duplicate_extra_and_tampered_members_rejected(self) -> None:
        prepared = self.build()
        with zipfile.ZipFile(self.archive) as archive:
            contents = [(name, archive.read(name)) for name in archive.namelist()]
        bad = self.root / "invalid.zip"
        for kind in ("missing", "duplicate", "extra", "tampered", "traversal", "wrongroot"):
            with self.subTest(kind=kind):
                entries = list(contents)
                if kind == "missing":
                    entries.pop()
                elif kind == "duplicate":
                    entries.append(entries[0])
                elif kind == "extra":
                    entries.append((self.output.name + "/producer.py", b"not permitted"))
                elif kind == "tampered":
                    entries[0] = (entries[0][0], entries[0][1] + b"tampered")
                elif kind == "traversal":
                    entries[0] = (self.output.name + "/../escape.txt", entries[0][1])
                else:
                    entries[0] = ("wrongroot/" + entries[0][0].split("/", 1)[1], entries[0][1])
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore", UserWarning)
                    with zipfile.ZipFile(bad, "w") as archive:
                        for name, data in entries:
                            archive.writestr(name, data)
                with self.assertRaises(B.ValidationError):
                    B.validate_zip(bad, self.output, prepared)

    def test_zip_crc_corruption_detected(self) -> None:
        prepared = self.build()
        # A stored archive gives a deterministic byte flip without a decompressor error.
        bad = self.root / "crc-bad.zip"
        with zipfile.ZipFile(bad, "w", compression=zipfile.ZIP_STORED) as archive:
            for name in sorted(prepared.expected_files):
                archive.writestr(f"{self.output.name}/{name}", (self.output / name).read_bytes())
        with zipfile.ZipFile(bad) as archive:
            member = archive.infolist()[0]
        contents = bytearray(bad.read_bytes())
        offset = member.header_offset
        filename_length, extra_length = struct.unpack_from("<HH", contents, offset + 26)
        contents[offset + 30 + filename_length + extra_length] ^= 1
        bad.write_bytes(contents)
        with self.assertRaisesRegex(B.ValidationError, "CRC|ZIP"):
            B.validate_zip(bad, self.output, prepared)

    def test_validate_only_reads_inputs_and_existing_package_without_writes(self) -> None:
        prepared = self.build()
        paths = [self.source_path, self.curation_path, self.archive] + [self.output / p for p in prepared.expected_files]
        before = {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in paths}
        with mock.patch("sys.stdout", new_callable=io.StringIO):
            self.assertEqual(B.main(["--source", str(self.source_path), "--curation", str(self.curation_path),
                                     "--output", str(self.output), "--zip", str(self.archive), "--validate-only"]), 0)
            self.assertEqual(B.main(["--source", str(self.source_path), "--curation", str(self.curation_path),
                                     "--validate-only"]), 0)
        self.assertEqual(before, {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in paths})

    def test_cli_rejects_incomplete_arguments_and_missing_inputs(self) -> None:
        with mock.patch("sys.stderr", new_callable=io.StringIO):
            self.assertEqual(B.main([]), 2)
            self.assertEqual(B.main(["--source", str(self.source_path), "--curation", str(self.curation_path),
                                     "--validate-only"]), 2)
            self.assertEqual(B.main(["--self-test", "--source", str(self.source_path)]), 2)

    def test_duplicate_json_keys_and_nonfinite_values_rejected(self) -> None:
        for value in (b'{"owner":"Berkay","owner":"Other"}', b'{"value":NaN}'):
            with self.subTest(value=value):
                with self.assertRaises(B.ValidationError):
                    B.parse_json(value, "synthetic input")


if __name__ == "__main__":
    unittest.main(verbosity=2)
