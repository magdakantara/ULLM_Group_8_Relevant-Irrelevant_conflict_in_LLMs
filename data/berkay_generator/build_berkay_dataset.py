#!/usr/bin/env python3
"""Build or validate the local, pending-review CUAD conflict data package.

Use Python's isolated mode when reading downloaded input:
    python3.13 -I -B build_berkay_dataset.py --source PATH --curation PATH \
        --output FRESH_DIRECTORY --zip FRESH_ARCHIVE.zip
    python3.13 -I -B build_berkay_dataset.py --source PATH --curation PATH \
        --validate-only [--output EXISTING_DIRECTORY --zip EXISTING_ARCHIVE.zip]
    python3.13 -I -B build_berkay_dataset.py --self-test

No networking, model inference, source editing, or legal/semantic approval occurs.
Offsets are zero-based Python Unicode character offsets; end offsets are exclusive.
"""

from __future__ import annotations

import argparse
import binascii
from collections import Counter
from dataclasses import dataclass
from datetime import date
import hashlib
import json
import os
from pathlib import Path
import re
import runpy
import stat
import sys
import tempfile
from typing import Any
import unittest
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COLUMNS = (
    "owner", "original_contract_id", "contract_type", "clause_category", "question",
    "source_A", "source_B", "exactly_what_changed", "case_label", "answer_supported_by_A",
    "answer_supported_by_B", "expected_behaviour", "reviewer", "review_status",
)
WORKBOOK = "contract_conflict_berkay_deliverable.xlsx"
SHEET = "Shared deliverable"
LABELS = ("agreement", "relevant_conflict", "irrelevant_conflict")
ALLOWED_CATEGORIES = {
    "Effective Date", "Expiration Date", "Renewal Term", "Notice Period To Terminate Renewal",
    "Governing Law", "Termination For Convenience", "Cap On Liability", "Warranty Duration",
}
ANSWER_BEHAVIOUR = "Give the supported common answer; do not refuse because of an unrelated document difference."
CONFLICT_BEHAVIOUR = "Report both source-specific answers and their difference; do not assume a governing version."
BASE_FILES = {WORKBOOK, "examples.jsonl", "curation.json", "README.txt"}
PROVENANCE_URL = "https://www.atticusprojectai.org/cuad"
REPOSITORY_URL = "https://github.com/TheAtticusProject/cuad"
DISCLAIMER = (
    "Automated checks establish mechanical preservation and consistency only; "
    "they are not semantic, legal, or human approval. No inference was performed."
)


class ValidationError(ValueError):
    """An input or generated package failed a fail-closed mechanical check."""


def require(condition: bool, message: str) -> None:
    """Raise a descriptive validation error, including under python -O."""
    if not condition:
        raise ValidationError(message)


def fields(value: Any, names: str, where: str) -> dict:
    """Require a JSON object and named fields, without adding schema fields."""
    require(isinstance(value, dict), f"{where}: expected an object")
    missing = set(names.split()) - value.keys()
    require(not missing, f"{where}: missing fields {sorted(missing)}")
    return value


def text(value: Any, where: str, *, empty: bool = False) -> str:
    """Require a UTF-8-encodable string without changing its contents."""
    require(isinstance(value, str), f"{where}: expected text")
    require(empty or bool(value.strip()), f"{where}: empty text")
    try:
        value.encode("utf-8")
    except UnicodeError as exc:
        raise ValidationError(f"{where}: invalid Unicode") from exc
    return value


def integer(value: Any, where: str) -> int:
    """Require an integer offset/count, not a bool or a coerced string."""
    require(type(value) is int and value >= 0, f"{where}: expected a nonnegative integer")
    return value


def digest(data: bytes) -> str:
    """Return the lowercase SHA-256 of exact bytes."""
    return hashlib.sha256(data).hexdigest()


def check_hash(expected: Any, data: bytes, where: str) -> None:
    """Check a supplied SHA-256 without normalizing the hashed data."""
    require(isinstance(expected, str) and re.fullmatch(r"[0-9a-fA-F]{64}", expected) is not None,
            f"{where}: invalid SHA-256")
    require(expected.lower() == digest(data), f"{where}: stale SHA-256")


def portable_name(value: Any, where: str) -> str:
    """Accept portable ASCII names, not paths, URL metacharacters, or device names."""
    value = text(value, where)
    require(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", value) is not None
            and not value.endswith(".") and len(value.encode("utf-8")) <= 200,
            f"{where}: use a portable ASCII basename (letters, digits, dot, underscore, hyphen)")
    stem = value.split(".")[0].upper()
    reserved = {"CON", "PRN", "AUX", "NUL", "CLOCK$", "CONIN$", "CONOUT$"}
    reserved.update(f"{prefix}{i}" for prefix in ("COM", "LPT") for i in range(1, 10))
    require(stem not in reserved, f"{where}: reserved Windows name")
    return value


def portable_path(value: str) -> None:
    """Check every component of a nonempty relative POSIX file path."""
    require(isinstance(value, str) and bool(value), "Empty/non-text package path")
    for part in value.split("/"):
        portable_name(part, f"package path {value!r}")


def parse_json(raw: bytes, where: str) -> Any:
    """Parse UTF-8 JSON, rejecting ambiguous duplicate keys and nonfinite values."""
    def unique_pairs(pairs: list[tuple[str, Any]]) -> dict:
        result = {}
        for key, value in pairs:
            require(key not in result, f"{where}: duplicate JSON key {key!r}")
            result[key] = value
        return result

    def invalid_constant(value: str) -> None:
        raise ValidationError(f"{where}: nonfinite JSON constant {value}")

    try:
        return json.loads(raw.decode("utf-8"), object_pairs_hook=unique_pairs,
                          parse_constant=invalid_constant)
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise ValidationError(f"{where}: invalid UTF-8 JSON: {exc}") from exc


def json_bytes(value: Any) -> bytes:
    """Serialize metadata; document texts never go through this serializer."""
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")


def validate_span(span: Any, context: str, where: str) -> dict:
    """Validate an exact, nonempty evidence quote in raw context."""
    fields(span, "start end text", where)
    start, end = integer(span["start"], where), integer(span["end"], where)
    quote = text(span["text"], where)
    require(start < end <= len(context), f"{where}: span out of range")
    require(context[start:end] == quote, f"{where}: stale offset or wrong evidence text")
    return {"start": start, "end": end, "text": quote}


def intersects(first: dict, second: dict) -> bool:
    """Test overlap of half-open intervals."""
    return first["start"] < second["end"] and second["start"] < first["end"]


def apply_patches(context: str, patches: list[dict]) -> str:
    """Apply validated edits independently to A, in descending original offsets."""
    result = context
    for patch in sorted(patches, key=lambda item: item["start"], reverse=True):
        require(result[patch["start"]:patch["end"]] == patch["before"], "Stale patch during reconstruction")
        result = result[:patch["start"]] + patch["after"] + result[patch["end"]:]
    return result


def remap_evidence(context: str, evidence: list[dict], patches: list[dict]) -> list[dict]:
    """Shift untouched spans and reconstruct edits wholly contained within a span."""
    modified = apply_patches(context, patches)
    mapped = []
    for span in evidence:
        before = [p for p in patches if p["end"] <= span["start"]]
        inside = [p for p in patches if span["start"] <= p["start"] and p["end"] <= span["end"]]
        require(all(not intersects(p, span) or p in inside for p in patches),
                "Patch crosses an evidence boundary")
        shift = sum(len(p["after"]) - len(p["before"]) for p in before)
        change = sum(len(p["after"]) - len(p["before"]) for p in inside)
        local = [{**p, "start": p["start"] - span["start"], "end": p["end"] - span["start"]}
                 for p in inside]
        mapped.append(validate_span({
            "start": span["start"] + shift, "end": span["end"] + shift + change,
            "text": apply_patches(span["text"], local),
        }, modified, "remapped evidence"))
    return mapped


def validate_patches(context: str, variant: dict, evidence: dict[str, list[dict]], target: str) -> list[dict]:
    """Require nonoverlapping real edits within the target's disjoint evidence."""
    fields(variant, "affected_question semantic_change patches", "variant")
    require(variant["affected_question"] == target, f"Variant must affect {target}")
    text(variant["semantic_change"], "semantic_change")
    patches = variant["patches"]
    require(isinstance(patches, list) and bool(patches), "Variant needs nonempty patches")
    result = []
    for patch in patches:
        fields(patch, "start end before after", "patch")
        start, end = integer(patch["start"], "patch start"), integer(patch["end"], "patch end")
        original = text(patch["before"], "patch before")
        replacement = text(patch["after"], "patch after", empty=True)
        require(start < end <= len(context), "Patch out of range or empty-before")
        require(context[start:end] == original, "Stale patch before string or offset")
        require(original != replacement, "No-op patch")
        require(any(s["start"] <= start < end <= s["end"] for s in evidence[target]),
                "Patch must be wholly inside target evidence; boundary crossing is forbidden")
        require(not any(intersects(patch, s) for qid, spans in evidence.items() if qid != target for s in spans),
                "Patch touches non-target evidence")
        result.append({"start": start, "end": end, "before": original, "after": replacement})
    result.sort(key=lambda item: item["start"])
    require(all(a["end"] <= b["start"] for a, b in zip(result, result[1:])), "Overlapping patches")
    require(apply_patches(context, result) != context, "Variant collectively makes no change")
    return result


def native_evidence(question: dict, paragraph: dict, context: str) -> list[dict]:
    """Bind exact curated subspans to the selected native QA's separate intervals.

    Fail closed on unreliable native offsets/text in this selected batch. Never
    normalize or concatenate native answers. Interval membership is mechanical
    provenance, not semantic validation.
    """
    annotation_id = text(question["annotation_id"], "annotation_id")
    require(annotation_id.endswith("__" + question["category"]), "Annotation/category mismatch")
    qas = paragraph.get("qas")
    require(isinstance(qas, list), "Source paragraph lacks qas")
    matches = [qa for qa in qas if isinstance(qa, dict) and qa.get("id") == annotation_id]
    require(len(matches) == 1, "Missing or ambiguous annotation_id in selected document")
    qa = matches[0]
    require(qa.get("is_impossible") is False, "Selected native annotation has no answer")
    answers = qa.get("answers")
    require(isinstance(answers, list) and bool(answers), "Selected annotation has no native spans")
    intervals = []
    for answer in answers:
        fields(answer, "answer_start text", "native answer")
        start = integer(answer["answer_start"], "native answer_start")
        end = start + len(text(answer["text"], "native answer text"))
        require(start < end <= len(context), "Native annotation interval out of range")
        require(context[start:end] == answer["text"], "Native answer text/offset does not exactly match raw context")
        intervals.append((start, end))
    spans = question["evidence"]
    require(isinstance(spans, list) and bool(spans), "Missing evidence spans")
    result = [validate_span(span, context, "original evidence") for span in spans]
    for span in result:
        require(any(start <= span["start"] < span["end"] <= end for start, end in intervals),
                "Evidence is not a subspan of the selected native annotation")
    return result


def change_description(variant: dict | None, patches: list[dict]) -> str:
    """Describe only real offsets/edits plus the curator's explicitly labeled rationale."""
    if variant is None:
        return "No edits. B0 is byte-for-byte identical to A."
    edits = "; ".join(
        f"A[{p['start']}:{p['end']}] {json.dumps(p['before'], ensure_ascii=False)} -> "
        f"{json.dumps(p['after'], ensure_ascii=False)}" for p in patches
    )
    return f"{edits}. Curated semantic change (pending review): {variant['semantic_change']}"


@dataclass
class Prepared:
    """One validated input snapshot and its canonical rows/documents."""

    curation: dict
    curation_raw: bytes
    source_hash: str
    rows: list[dict]
    documents: dict[str, bytes]

    @property
    def expected_files(self) -> set[str]:
        return BASE_FILES | {"validation.json"} | self.documents.keys()


def prepare(source_raw: bytes, curation_raw: bytes, source_basename: str) -> Prepared:
    """Validate the agreed schema and derive all six rows per original contract."""
    source = fields(parse_json(source_raw, "source"), "version data", "source")
    curation = fields(parse_json(curation_raw, "curation"),
                      "schema_version created_date source owner reviewer review_status expected_contract_count contracts",
                      "curation")
    require(curation["schema_version"] == "1.0", "Unsupported curation schema_version")
    try:
        created = text(curation["created_date"], "created_date")
        require(date.fromisoformat(created).isoformat() == created, "created_date must be YYYY-MM-DD")
    except ValueError as exc:
        raise ValidationError("created_date must be YYYY-MM-DD") from exc
    provenance = fields(curation["source"], "filename sha256 version", "curation source")
    require(portable_name(provenance["filename"], "source filename") == source_basename,
            "Source filename does not match the supplied basename")
    require(text(provenance["version"], "source version") == source["version"], "Source version mismatch")
    check_hash(provenance["sha256"], source_raw, "source")
    require(curation["owner"] == "Berkay" and curation["reviewer"] == "Engin",
            "This package requires owner Berkay and reviewer Engin")
    require(curation["review_status"] == "pending", "Human review must remain pending")
    count = integer(curation["expected_contract_count"], "expected_contract_count")
    contracts = curation["contracts"]
    require(count > 0 and isinstance(contracts, list) and len(contracts) == count,
            "Contract count differs from expected_contract_count")
    data = source["data"]
    require(isinstance(data, list), "Source data must be a list")
    title_counts = Counter(doc.get("title") for doc in data if isinstance(doc, dict) and isinstance(doc.get("title"), str))
    seen_ids, seen_titles, seen_indices, seen_contexts = set(), set(), set(), set()
    rows, documents = [], {}
    for contract in contracts:
        fields(contract, "id original_contract_id contract_type document_index paragraph_index context_sha256 "
                         "questions variants independence_note", "contract")
        key = portable_name(contract["id"], "contract id")
        title = text(contract["original_contract_id"], "original_contract_id")
        contract_type = text(contract["contract_type"], "contract_type")
        index = integer(contract["document_index"], "document_index")
        paragraph_index = integer(contract["paragraph_index"], "paragraph_index")
        require(key.casefold() not in seen_ids and title not in seen_titles and index not in seen_indices,
                "Duplicate contract id, original title, or document index")
        require(index < len(data), "document_index out of range")
        doc = fields(data[index], "title paragraphs", "source document")
        require(doc["title"] == title and title_counts[title] == 1, "Original title/index mismatch or ambiguous title")
        require(isinstance(doc["paragraphs"], list) and len(doc["paragraphs"]) == 1 and paragraph_index == 0,
                "Schema 1.0 requires exactly one full context per document and paragraph_index 0")
        paragraph = fields(doc["paragraphs"][0], "context qas", "source paragraph")
        context = text(paragraph["context"], "raw context")
        context_bytes = context.encode("utf-8")
        check_hash(contract["context_sha256"], context_bytes, "context")
        require(digest(context_bytes) not in seen_contexts, "Duplicate full contract context")
        seen_ids.add(key.casefold())
        seen_titles.add(title)
        seen_indices.add(index)
        seen_contexts.add(digest(context_bytes))
        text(contract["independence_note"], "independence_note")
        notes = contract.get("review_notes", [])
        require(notes is None or isinstance(notes, list) and all(isinstance(n, str) for n in notes),
                "review_notes must be a list of strings when supplied")
        questions = contract["questions"]
        require(isinstance(questions, list) and len(questions) == 2, "Each contract needs exactly Q1 and Q2")
        by_question, evidence = {}, {}
        for question in questions:
            fields(question, "id category text answer_original answer_modified annotation_id evidence", "question")
            qid = question["id"]
            require(isinstance(qid, str) and qid in {"Q1", "Q2"} and qid not in by_question,
                    "Question ids must be unique Q1 and Q2")
            for name in ("category", "text", "answer_original", "answer_modified"):
                text(question[name], f"{qid}.{name}")
            require(question["category"] in ALLOWED_CATEGORIES, "Unsupported clause category for this study")
            require(question["answer_original"].strip() != question["answer_modified"].strip(),
                    "Inconsistent answers: target original and modified answers must differ")
            if question.get("scope_notes") is not None:
                text(question["scope_notes"], "scope_notes", empty=True)
            by_question[qid] = question
            evidence[qid] = native_evidence(question, paragraph, context)
        all_spans = sorted([s for spans in evidence.values() for s in spans], key=lambda s: s["start"])
        require(all(a["end"] <= b["start"] for a, b in zip(all_spans, all_spans[1:])),
                "Overlapping evidence fails closed (within or across questions)")
        variants = fields(contract["variants"], "B1 B2", "variants")
        require(set(variants) == {"B1", "B2"}, "Only independently derived B1 and B2 variants are supported")
        prefix = f"documents/{key}"
        documents[f"{prefix}/A.txt"] = context_bytes
        for variant_id in ("B0", "B1", "B2"):
            variant = variants.get(variant_id)
            target = {"B1": "Q1", "B2": "Q2"}.get(variant_id)
            patches = validate_patches(context, variant, evidence, target) if variant_id != "B0" else []
            version_text = apply_patches(context, patches)
            documents[f"{prefix}/{variant_id}.txt"] = version_text.encode("utf-8")
            for qid in ("Q1", "Q2"):
                question = by_question[qid]
                changed = qid == target
                mapped = remap_evidence(context, evidence[qid], patches)
                if not changed:
                    require([s["text"] for s in mapped] == [s["text"] for s in evidence[qid]],
                            "Non-target evidence changed")
                # Answers are AI-assisted curated strings, not native extraction
                # labels or necessarily literal quotes. Only selection/inequality
                # is mechanically checked; semantic support remains pending.
                label = "agreement" if variant_id == "B0" else "relevant_conflict" if changed else "irrelevant_conflict"
                answer_b = question["answer_modified"] if changed else question["answer_original"]
                rows.append({
                    "owner": curation["owner"], "original_contract_id": title, "contract_type": contract_type,
                    "clause_category": question["category"], "question": question["text"],
                    "source_A": f"{prefix}/A.txt", "source_B": f"{prefix}/{variant_id}.txt",
                    "exactly_what_changed": change_description(variant, patches), "case_label": label,
                    "answer_supported_by_A": question["answer_original"], "answer_supported_by_B": answer_b,
                    "expected_behaviour": CONFLICT_BEHAVIOUR if changed else ANSWER_BEHAVIOUR,
                    "reviewer": curation["reviewer"], "review_status": "pending",
                    "example_id": f"{key}-{variant_id}-{qid}", "contract_key": key,
                    "question_id": qid, "variant": variant_id,
                    "evidence": {"A": evidence[qid], "B": mapped},
                    "scope_notes": question.get("scope_notes") or "",
                    "provenance": {
                        "source_filename": provenance["filename"], "source_sha256": provenance["sha256"].lower(),
                        "source_version": source["version"], "document_index": index,
                        "paragraph_index": paragraph_index, "annotation_id": question["annotation_id"],
                        "context_sha256": digest(context_bytes), "source_url": PROVENANCE_URL,
                        "answer_origin": "Curator-provided; variant-selected; semantic review pending",
                    },
                })
    prepared = Prepared(curation, curation_raw, digest(source_raw), rows, documents)
    validate_rows(rows, count)
    for row in rows:
        for column in COLUMNS:
            spreadsheet_text(row[column])
    require(len({p.casefold() for p in prepared.expected_files}) == len(prepared.expected_files),
            "Case-insensitive package path collision")
    for path in prepared.expected_files:
        portable_path(path)
    return prepared


def validate_rows(rows: list[dict], count: int) -> None:
    """Check counts, matched labels, answer identity, and pending review status."""
    require(len(rows) == 6 * count, "Expected six rows per contract")
    groups = {}
    seen = set()
    for row in rows:
        fields(row, " ".join(COLUMNS) + " example_id contract_key question_id variant evidence provenance", "example")
        require(row["example_id"] not in seen, "Duplicate example_id")
        seen.add(row["example_id"])
        require(row["reviewer"] == "Engin" and row["review_status"] == "pending" and row["owner"] == "Berkay",
                "All examples require Berkay / Engin / pending")
        variant, qid = row["variant"], row["question_id"]
        require(variant in {"B0", "B1", "B2"} and qid in {"Q1", "Q2"}, "Invalid question/variant")
        relevant = (variant, qid) in {("B1", "Q1"), ("B2", "Q2")}
        label = "agreement" if variant == "B0" else "relevant_conflict" if relevant else "irrelevant_conflict"
        require(row["case_label"] == label, "Incorrect case label")
        require(row["expected_behaviour"] == (CONFLICT_BEHAVIOUR if relevant else ANSWER_BEHAVIOUR),
                "Incorrect expected behaviour")
        require((row["answer_supported_by_A"] != row["answer_supported_by_B"]) == relevant,
                "Inconsistent source answers")
        groups.setdefault(row["contract_key"], []).append(row)
    require(len(groups) == count, "Distinct contract count mismatch")
    for group in groups.values():
        require(Counter(r["case_label"] for r in group) == Counter({label: 2 for label in LABELS}),
                "Each contract must have two examples per label")
        require({(r["variant"], r["question_id"]) for r in group}
                == {(v, q) for v in ("B0", "B1", "B2") for q in ("Q1", "Q2")},
                "Missing or duplicated matched question/variant")


def spreadsheet_text(value: str) -> None:
    """Reject Excel truncation and XML-invalid text before assigning any cell."""
    text(value, "spreadsheet cell", empty=True)
    require(len(value) <= 32767 and len(value.encode("utf-16-le")) // 2 <= 32767,
            "Spreadsheet cell exceeds 32767 characters/code units")
    require(ILLEGAL_CHARACTERS_RE.search(value) is None and "￾" not in value and "￿" not in value,
            "Spreadsheet cell contains XML-invalid characters")


def write_workbook(path: Path, rows: list[dict]) -> None:
    """Write exactly the shared 14 columns, always as literal text, with local links."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    for row_number, values in enumerate([COLUMNS] + [tuple(row[c] for c in COLUMNS) for row in rows], 1):
        for column_number, value in enumerate(values, 1):
            spreadsheet_text(value)
            cell = sheet.cell(row_number, column_number, value)
            cell.data_type = "s"  # Also overrides formula and Excel error-string detection.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_number == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="234E52")
            elif COLUMNS[column_number - 1] in {"source_A", "source_B"}:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
    widths = (14, 48, 25, 24, 58, 37, 37, 80, 24, 44, 44, 22, 14, 17)
    for number, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    workbook.save(path)
    workbook.close()


def readme(prepared: Prepared) -> bytes:
    """Describe research limitations and portable package use in English."""
    count = prepared.curation["expected_contract_count"]
    return (f"""Berkay local contract-conflict data package

SYNTHETIC RESEARCH ONLY. Not legal advice or an authoritative contract revision.
AI-assisted curation; owner Berkay, reviewer Engin. All human reviews are pending.
{DISCLAIMER}

Provenance: CUAD, {PROVENANCE_URL}
Original dataset repository: {REPOSITORY_URL}
Source: {prepared.curation['source']['filename']} (version {prepared.curation['source']['version']})
Source SHA-256: {prepared.source_hash}
Original CUAD authors/annotations are not authors or approvers of these synthetic edits.

This is a purposive sample, not a random or representative sample. There are
{count} distinct original contracts, {4 * count} document files, and {6 * count} examples
({2 * count} per label). Matched examples from the same contract are correlated;
do not treat them as independent observations or split them across train/test.
Using these examples for prompt design disqualifies calling them an untouched
final test set. No model run, inference, evaluation score, or performance claim
is included. No source-priority, recency, or authoritative-version rule applies.

A contains the original full contract context, exactly encoded as UTF-8 without
normalization or an added newline. B0 is byte-identical to A. B1 and B2 are each
independently derived from A, not sequentially from each other. Offsets are
zero-based Unicode character positions with exclusive ends, not byte positions.
B1 changes Q1 evidence; B2 changes Q2 evidence. For each question, agreement and
irrelevant_conflict expect ANSWER using the common answer. relevant_conflict
expects abstention from choosing one source: report both source-specific answers
and explain the disagreement. PRESENT_BOTH is therefore acceptable behaviour, not
an unsupported single answer. The expected_behaviour column uses the team's shared
descriptive wording. Labels and answers are AI-assisted curator-provided drafts
pending human semantic/legal review; no independent human adjudication is claimed.
Evidence preservation alone does not establish whole-contract semantic independence.
Native QA answers are separate provenance intervals; exact curated evidence may
be a strict subspan. Native answer strings are not normalized or concatenated.

Files:
- {WORKBOOK}: sheet '{SHEET}', the shared 14 columns. All cells are literal text.
- examples.jsonl: the same ordered rows, plus evidence offsets and provenance.
- curation.json: exact input manifest bytes; includes edits and rationale.
- documents/<contract-id>/A.txt, B0.txt, B1.txt, B2.txt: full source texts.
- validation.json: automated integrity checks and hashes, NOT human approval.

Keep this entire directory together. Workbook links are relative POSIX paths
and resolve when the ZIP is extracted with its top-level dataset folder intact.
Some spreadsheet viewers require permission before opening local file links.
The producer script, original source JSON, and any unrelated files are not bundled.
validation.json deliberately excludes its own hash and any archive hash to avoid
self-referential cycles. ZIP names, CRCs and bytes are checked after finalization
by the builder and can be checked again with --validate-only and --zip.
""").encode("utf-8")


def inventory(root: Path, expected: set[str]) -> None:
    """Inspect only the specified package root; reject symlinks and unexpected files."""
    require(root.is_dir() and not root.is_symlink(), "Package directory missing or symlinked")
    expected_dirs = {str(Path(p).parent).replace(os.sep, "/") for p in expected if "/" in p} | {"documents"}
    actual_files, actual_dirs = set(), set()
    for directory, dirs, files in os.walk(root, followlinks=False):
        for name in dirs + files:
            path = Path(directory) / name
            relative = path.relative_to(root).as_posix()
            portable_path(relative)
            require(not path.is_symlink(), f"Symlink in package: {relative}")
            if name in dirs:
                require(relative in expected_dirs, f"Unexpected directory: {relative}")
                actual_dirs.add(relative)
            else:
                require(path.is_file(), f"Non-regular file: {relative}")
                actual_files.add(relative)
    require(actual_dirs == expected_dirs, "Missing package directories")
    require(actual_files == expected, f"Package file mismatch: missing={sorted(expected - actual_files)}, "
            f"unexpected={sorted(actual_files - expected)}")


def validation_report(root: Path, prepared: Prepared) -> dict:
    """Hash all payloads except this report, without archive/self-hash cycles."""
    return {
        "schema_version": "1.0", "automated_validation": "passed", "disclaimer": DISCLAIMER,
        "human_review_status": "pending", "semantic_or_legal_approval": False, "inference_performed": False,
        "contract_count": prepared.curation["expected_contract_count"], "row_count": len(prepared.rows),
        "document_count": len(prepared.documents), "label_counts": dict(Counter(r["case_label"] for r in prepared.rows)),
        "source": prepared.curation["source"], "curation_sha256": digest(prepared.curation_raw),
        "checks": ["source and context hashes", "title and document index", "native annotation membership",
                   "exact evidence and remapped offsets", "independent variant reconstruction", "A/B0 byte identity",
                   "target-only nonoverlapping patches", "six matched rows and two of each label per contract",
                   "portable paths and exact file inventory", "XLSX/JSONL values, literal cell types and links",
                   "all human review pending"],
        "expected_files": sorted(prepared.expected_files),
        "sha256": {name: digest((root / name).read_bytes())
                   for name in sorted(prepared.expected_files - {"validation.json"})},
        "hash_exclusions": ["validation.json (self-reference)", "ZIP archive (created and verified last)"],
        "archive_validation": "Not attested by this pre-archive report. Builder checks unique names, CRCs and exact bytes "
                              "after finalization; --validate-only --zip rechecks them without modifying this report.",
    }


def validate_package(root: Path, prepared: Prepared, *, with_report: bool = True) -> None:
    """Reopen and compare all package contents with the validated input snapshot."""
    expected = prepared.expected_files if with_report else prepared.expected_files - {"validation.json"}
    inventory(root, expected)
    require((root / "curation.json").read_bytes() == prepared.curation_raw, "Curation copy changed")
    require((root / "README.txt").read_bytes() == readme(prepared), "README changed")
    for name, expected_bytes in prepared.documents.items():
        require((root / name).read_bytes() == expected_bytes, f"Document missing, tampered, or incorrectly derived: {name}")
    for contract in prepared.curation["contracts"]:
        prefix = root / "documents" / contract["id"]
        require((prefix / "A.txt").read_bytes() == (prefix / "B0.txt").read_bytes(), "A and B0 differ")
    lines = (root / "examples.jsonl").read_bytes().splitlines()
    require(all(lines), "Blank JSONL line")
    rows = [parse_json(line, f"examples.jsonl line {i}") for i, line in enumerate(lines, 1)]
    validate_rows(rows, prepared.curation["expected_contract_count"])
    require(rows == prepared.rows, "JSONL differs from canonical rows/answers/evidence/provenance")
    for row in rows:
        for column, version in (("source_A", "A"), ("source_B", "B")):
            portable_path(row[column])
            require(row[column] in prepared.documents, "Source reference is not an expected document")
            version_text = (root / row[column]).read_bytes().decode("utf-8")
            for span in row["evidence"][version]:
                validate_span(span, version_text, f"{row['example_id']} {version} quote")
    workbook = load_workbook(root / WORKBOOK, data_only=False, keep_links=False)
    try:
        require(workbook.sheetnames == [SHEET], "Unexpected workbook sheets")
        sheet = workbook[SHEET]
        require(sheet.max_row == len(rows) + 1 and sheet.max_column == len(COLUMNS), "Workbook dimensions mismatch")
        expected_values = [COLUMNS] + [tuple(row[c] for c in COLUMNS) for row in rows]
        for cells, values in zip(sheet.iter_rows(), expected_values):
            for cell, value in zip(cells, values):
                require(cell.data_type == "s" and cell.value == value, f"Formula, truncation or workbook mismatch: {cell.coordinate}")
                is_link = cell.row > 1 and COLUMNS[cell.column - 1] in {"source_A", "source_B"}
                require((cell.hyperlink.target if cell.hyperlink else None) == (value if is_link else None),
                        f"Incorrect or unexpected hyperlink: {cell.coordinate}")
        require(sheet.freeze_panes == "A2" and sheet.auto_filter.ref == sheet.dimensions,
                "Workbook header freeze/filter mismatch")
    finally:
        workbook.close()
    if with_report:
        require(parse_json((root / "validation.json").read_bytes(), "validation.json") == validation_report(root, prepared),
                "Validation report or payload hashes changed")


def validate_zip(archive: Path, root: Path, prepared: Prepared) -> None:
    """Require the exact top-level folder, unique portable names, CRCs and payloads."""
    expected = {f"{root.name}/{name}": name for name in prepared.expected_files}
    try:
        with zipfile.ZipFile(archive, "r") as zipped:
            members = zipped.infolist()
            names = [item.filename for item in members]
            require(len(names) == len(set(names)) == len({n.casefold() for n in names}), "Duplicate ZIP member names")
            require(set(names) == set(expected), "ZIP file inventory or top-level folder mismatch")
            require(zipped.testzip() is None, "ZIP CRC check failed")
            for member in members:
                portable_path(member.filename)
                require(not member.is_dir() and not stat.S_ISLNK(member.external_attr >> 16), "ZIP contains directory/symlink member")
                contents = zipped.read(member)
                require(member.CRC == binascii.crc32(contents) & 0xFFFFFFFF, "ZIP CRC mismatch")
                require(member.file_size == len(contents), "ZIP file size mismatch")
                require(contents == (root / expected[member.filename]).read_bytes(), f"ZIP bytes differ: {member.filename}")
    except (zipfile.BadZipFile, EOFError, RuntimeError) as exc:
        raise ValidationError(f"Invalid ZIP: {exc}") from exc


def check_unchanged(source: Path, curation: Path, prepared: Prepared) -> None:
    """Detect source/manifest modifications during a build or validation."""
    require(digest(source.read_bytes()) == prepared.source_hash, "Source changed during operation")
    require(curation.read_bytes() == prepared.curation_raw, "Curation changed during operation")


def fresh_destination(path: Path) -> Path:
    """Require existing parents and a wholly fresh destination, including no dangling link."""
    portable_name(path.name, "destination basename")
    require(path.parent.is_dir(), f"Destination parent does not exist: {path.parent}")
    path = path.parent.resolve() / path.name
    require(not os.path.lexists(path), f"Destination already exists (fresh target required): {path}")
    return path


def build(source: Path, curation: Path, output: Path, archive: Path) -> Prepared:
    """Stage and validate a fresh directory, publish it, then create/verify ZIP last.

    Existing targets, including empty directories/files, are never overwritten.
    On a late ZIP failure the already validated output directory is retained.
    Only private temporary paths owned by this operation are cleaned up.
    """
    output, archive = fresh_destination(output), fresh_destination(archive)
    require(str(output).casefold() != str(archive).casefold() and not archive.is_relative_to(output),
            "ZIP must be outside the dataset directory without case-insensitive collisions")
    for input_path in (source, curation, Path(__file__)):
        require(not input_path.resolve().is_relative_to(output), "Input/producer cannot be inside output")
        require(input_path.resolve() != archive, "ZIP collides with an input/producer")
    prepared = prepare(source.read_bytes(), curation.read_bytes(), source.name)
    with tempfile.TemporaryDirectory(prefix=".berkay-build-", dir=output.parent) as temporary:
        staging = Path(temporary) / output.name
        staging.mkdir()
        for name, contents in prepared.documents.items():
            destination = staging / name
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(contents)
        (staging / "curation.json").write_bytes(prepared.curation_raw)
        (staging / "README.txt").write_bytes(readme(prepared))
        (staging / "examples.jsonl").write_bytes(b"".join(
            (json.dumps(row, ensure_ascii=False, allow_nan=False) + "\n").encode("utf-8") for row in prepared.rows
        ))
        write_workbook(staging / WORKBOOK, prepared.rows)
        validate_package(staging, prepared, with_report=False)
        (staging / "validation.json").write_bytes(json_bytes(validation_report(staging, prepared)))
        validate_package(staging, prepared)
        check_unchanged(source, curation, prepared)
        fresh_destination(archive)
        # An exclusive empty reservation avoids overwriting a pre-existing output.
        output.mkdir(exist_ok=False)
        try:
            staging.rename(output)
        except BaseException:
            output.rmdir()  # Only our own still-empty reservation, never a directory sweep.
            raise
    with tempfile.TemporaryDirectory(prefix=".berkay-zip-", dir=archive.parent) as temporary:
        zip_staging = Path(temporary) / "package.zip"
        with zipfile.ZipFile(zip_staging, "x", compression=zipfile.ZIP_DEFLATED) as zipped:
            for name in sorted(prepared.expected_files):
                zipped.write(output / name, arcname=f"{output.name}/{name}")
        validate_package(output, prepared)
        validate_zip(zip_staging, output, prepared)
        check_unchanged(source, curation, prepared)
        os.link(zip_staging, archive)  # Atomic no-clobber publication on the same filesystem.
    return prepared


def self_test() -> bool:
    """Run only the adjacent synthetic unittest suite; all fixtures use tempfiles."""
    test_path = Path(__file__).with_name("test_build_berkay_dataset.py")
    require(test_path.is_file(), "Synthetic self-test file is missing")
    namespace = runpy.run_path(str(test_path), init_globals={"BUILDER": sys.modules[__name__]})
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(namespace["BuilderTests"])
    return unittest.TextTestRunner(verbosity=2).run(suite).wasSuccessful()


def main(argv: list[str] | None = None) -> int:
    """CLI entry point. Validation-only never writes files."""
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source", type=Path, help="Original CUAD JSON (read-only)")
    parser.add_argument("--curation", type=Path, help="Agreed curation JSON (read-only)")
    parser.add_argument("--output", type=Path, help="Fresh output folder, or existing folder with --validate-only")
    parser.add_argument("--zip", dest="archive", type=Path, help="Fresh ZIP, or existing ZIP with --validate-only")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--validate-only", action="store_true")
    mode.add_argument("--self-test", action="store_true")
    args = parser.parse_args(argv)
    try:
        if args.self_test:
            require(not any((args.source, args.curation, args.output, args.archive)), "--self-test accepts no data paths")
            return 0 if self_test() else 1
        require(args.source is not None and args.curation is not None, "--source and --curation are required")
        require(sys.flags.isolated, "Use python3.13 -I to safely read downloaded JSON")
        if args.validate_only:
            require(args.archive is None or args.output is not None, "--zip validation also needs --output")
            prepared = prepare(args.source.read_bytes(), args.curation.read_bytes(), args.source.name)
            if args.output is not None:
                validate_package(args.output, prepared)
            if args.archive is not None:
                validate_zip(args.archive, args.output, prepared)
            check_unchanged(args.source, args.curation, prepared)
            print(f"Automated validation passed: {len(prepared.rows)} rows; no files written.")
        else:
            require(args.output is not None and args.archive is not None, "Build requires --output and --zip (fresh targets)")
            prepared = build(args.source, args.curation, args.output, args.archive)
            print(f"Built and verified {len(prepared.rows)} rows: {args.output}\nZIP names/CRCs/bytes verified: {args.archive}")
        print(DISCLAIMER)
        return 0
    except (ValidationError, OSError, zipfile.BadZipFile) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
